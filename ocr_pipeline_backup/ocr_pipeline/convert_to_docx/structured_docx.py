"""
Build Word tables from structured_tables.json (no pandas dependency).

Used by convert_to_docx worker. Form-agnostic: uses header_matrix when present.
"""

from __future__ import annotations

from typing import Any, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def _set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(10)


def _merge_cells_horizontal(table, row_idx: int, start_c: int, end_c: int) -> None:
    if end_c <= start_c:
        return
    cell_a = table.cell(row_idx, start_c)
    cell_b = table.cell(row_idx, end_c)
    cell_a.merge(cell_b)


def _merge_cells_vertical(table, col_idx: int, start_r: int, end_r: int) -> None:
    if end_r <= start_r:
        return
    cell_a = table.cell(start_r, col_idx)
    cell_b = table.cell(end_r, col_idx)
    cell_a.merge(cell_b)


def _lower_rows_have_labels(
    matrix: Sequence[Sequence[str]], row_idx: int, start_c: int, end_c: int
) -> bool:
    for r in range(row_idx + 1, len(matrix)):
        for c in range(start_c, end_c + 1):
            if c < len(matrix[r]) and str(matrix[r][c]).strip():
                return True
    return False


def _looks_like_group(label: str) -> bool:
    text = label.strip().lower()
    if not text:
        return False
    markers = ("trong đó", "trong do", "bao gồm", "bao gom", "chi tiết", "chi tiet")
    return any(m in text for m in markers) or len(text) > 40


def _apply_header_merges(table, matrix: list[list[str]]) -> None:
    n_rows = len(matrix)
    n_cols = max((len(r) for r in matrix), default=0)
    for r in range(n_rows):
        c = 0
        while c < n_cols:
            cell = matrix[r][c].strip() if c < len(matrix[r]) else ""
            if not cell:
                c += 1
                continue
            end = c
            while end + 1 < n_cols:
                nxt = matrix[r][end + 1].strip() if end + 1 < len(matrix[r]) else ""
                if nxt:
                    break
                end += 1
            if end > c and (_looks_like_group(cell) or _lower_rows_have_labels(matrix, r, c + 1, end)):
                _merge_cells_horizontal(table, r, c, end)
            c = end + 1

    for c in range(n_cols):
        r = 0
        while r < n_rows:
            cell = matrix[r][c].strip() if c < len(matrix[r]) else ""
            if not cell:
                r += 1
                continue
            end = r
            while end + 1 < n_rows:
                below = matrix[end + 1][c].strip() if c < len(matrix[end + 1]) else ""
                if below:
                    break
                end += 1
            if end > r:
                _merge_cells_vertical(table, c, r, end)
            r = end + 1


def add_structured_table(doc: Document, table_obj: dict[str, Any]) -> None:
    """Append one structured table object as a Word grid."""
    columns = [str(c) for c in table_obj.get("columns") or []]
    data = table_obj.get("data") or []
    matrix_raw = table_obj.get("header_matrix")
    matrix: list[list[str]] | None = None
    if isinstance(matrix_raw, list) and matrix_raw:
        matrix = [[str(c) if c is not None else "" for c in row] for row in matrix_raw]

    n_cols = len(columns)
    if matrix:
        n_cols = max(n_cols, max(len(r) for r in matrix))
    if not n_cols and data:
        n_cols = max(len(r) for r in data)
    if n_cols <= 0:
        return

    n_hdr = len(matrix) if matrix else (1 if columns else 0)
    n_body = len(data)
    n_rows = n_hdr + n_body
    if n_rows <= 0:
        return

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"

    if matrix:
        for r, row in enumerate(matrix):
            for c in range(n_cols):
                val = row[c] if c < len(row) else ""
                _set_cell_text(table.rows[r].cells[c], val, bold=True)
                table.rows[r].cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _apply_header_merges(table, matrix)
    elif columns:
        for c, name in enumerate(columns):
            _set_cell_text(table.rows[0].cells[c], name, bold=True)

    for r, row in enumerate(data):
        excel_r = n_hdr + r
        for c in range(n_cols):
            val = row[c] if c < len(row) else ""
            _set_cell_text(table.rows[excel_r].cells[c], str(val) if val is not None else "")

    doc.add_paragraph("")


def add_structured_tables_from_payload(doc: Document, payload: dict[str, Any]) -> int:
    tables = payload.get("tables") if isinstance(payload, dict) else None
    if not isinstance(tables, list):
        return 0
    count = 0
    for t in tables:
        if isinstance(t, dict):
            add_structured_table(doc, t)
            count += 1
    return count


def create_document_from_pages_and_tables(
    pages: list[dict[str, Any]],
    structured_payload: dict[str, Any] | None = None,
) -> Document:
    """
    Prose lines from page JSON, then structured tables (preferred).
    Falls back to crude page['tables'] only when structured_payload is absent.
    """
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    # East Asia font hint for Word on Windows
    try:
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    except Exception:
        pass

    for page_idx, page in enumerate(pages):
        if page_idx > 0:
            doc.add_page_break()
        for line in page.get("lines") or []:
            text = (line.get("line_text") or "").strip()
            if text:
                doc.add_paragraph(text)

        if structured_payload is None:
            for t in page.get("tables") or []:
                # Legacy crude reconstruct tables
                t_lines = t.get("lines") or []
                if not t_lines:
                    continue
                num_rows = len(t_lines)
                num_cols = max((len(line.get("cells") or []) for line in t_lines), default=0)
                if num_rows <= 0 or num_cols <= 0:
                    continue
                word_table = doc.add_table(rows=num_rows, cols=num_cols)
                word_table.style = "Table Grid"
                for r_idx, row in enumerate(t_lines):
                    cells = row.get("cells") or []
                    for c_idx, cell in enumerate(cells):
                        if c_idx < num_cols:
                            word_table.cell(r_idx, c_idx).text = cell.get("text") or ""
                doc.add_paragraph()

    if structured_payload is not None:
        doc.add_paragraph()
        add_structured_tables_from_payload(doc, structured_payload)

    return doc
