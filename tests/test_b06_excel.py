"""Tests for B06 hierarchical Excel export (merged parent/child headers)."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from pdf_extractor.b06_excel import is_b06_excel_candidate, write_b06_workbook
from pdf_extractor.table_layout import _B06_CANONICAL_HEADERS


class TestB06Excel(unittest.TestCase):
    def _sample_df(self) -> pd.DataFrame:
        rows = [
            [1, "I", "Số phải thu", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
            [1, "1", "Các khoản chủ động", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
            [2, "1.1", "  Thu nộp", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
            [3, "1.1.1", "    Án phí", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
        ]
        cols = ["Cấp"] + list(_B06_CANONICAL_HEADERS)
        return pd.DataFrame(rows, columns=cols)

    def test_detects_b06_candidate(self) -> None:
        self.assertTrue(is_b06_excel_candidate(self._sample_df()))

    def test_writes_merged_group_header_and_codes(self) -> None:
        df = self._sample_df()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            write_b06_workbook([df], path)
            wb = load_workbook(path)
            ws = wb.active
            # Cấp | Số TT | Tiêu chí | Tổng | group... → excel col 5 is start of group (offset=1)
            self.assertEqual(ws.cell(1, 2).value, "Số TT")
            self.assertIn("Trong đó Chấp hành viên", str(ws.cell(1, 5).value))
            self.assertEqual(ws.cell(2, 5).value, "Ủy thác THA")
            self.assertEqual(ws.cell(2, 6).value, "Trả đơn THA")
            self.assertEqual(ws.cell(3, 2).value, "A")
            self.assertEqual(ws.cell(3, 5).value, "2")
            self.assertEqual(ws.cell(3, 10).value, "7")
            # Group merge covers columns 5-8 on row 1
            ranges = {str(r) for r in ws.merged_cells.ranges}
            self.assertTrue(any("E1:H1" in r or "E1:H1" == r for r in ranges) or any(
                r.startswith("E1:") and "H1" in r for r in ranges
            ))
            # Body starts at row 4; parent row bold (col2=STT with Cấp offset)
            self.assertEqual(ws.cell(4, 2).value, "I")
            self.assertTrue(ws.cell(4, 2).font.bold)
            self.assertEqual(ws.cell(7, 2).value, "1.1.1")
            self.assertEqual(ws.cell(4, 3).value, "Số phải thu")


if __name__ == "__main__":
    unittest.main()
