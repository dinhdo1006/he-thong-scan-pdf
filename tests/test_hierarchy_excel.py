"""Tests for generic hierarchical Excel export (multi-row header merge)."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from pdf_extractor.hierarchy_excel import drop_semantic_helpers, write_hierarchy_workbook
from pdf_extractor.table_layout import repair_form_table


class TestHierarchyExcel(unittest.TestCase):
    def _sample_df(self) -> pd.DataFrame:
        df = pd.DataFrame(
            [
                ["", "", "", "Ủy thác THA", "Trả đơn THA", "Đình chỉ THA", "Miễn, giảm THA", "", ""],
                ["A", "B", "1", "2", "3", "4", "5", "6", "7"],
                ["I", "số phải thu", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
                ["1", "Các khoản chủ động", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
                ["1.1", "Thu nộp", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
                ["1.1.1", "Án phí", "10.620.000", "10.370.000", "", "", "", "250.000", "250.000"],
            ],
            columns=[
                "Số TT",
                "Tiêu chí trong quyết định thi hành án",
                "Tổng số tiền, giá trị tài sản phải thi hành",
                "Trong đó Chấp hành viên đã giải quyết bằng các biện pháp",
                "",
                "",
                "",
                "Số thực thu thi hành án",
                "Số đã nộp Nhà nước và chi trả đương sự",
            ],
        )
        return repair_form_table(df)

    def test_repair_stores_header_matrix(self) -> None:
        df = self._sample_df()
        matrix = df.attrs.get("header_matrix")
        self.assertIsInstance(matrix, list)
        self.assertGreaterEqual(len(matrix), 2)
        # Ủy thác should sit under the group column positionally.
        flat = [str(c) for c in df.columns]
        self.assertTrue(any("Ủy thác" in c for c in flat))

    def test_drop_helpers(self) -> None:
        df = self._sample_df()
        df["STT_valid"] = True
        df["Sum_check"] = "OK"
        out = drop_semantic_helpers(df)
        names = [str(c).lower() for c in out.columns]
        self.assertNotIn("stt_valid", names)
        self.assertNotIn("sum_check", names)

    def test_writes_merged_group_from_matrix(self) -> None:
        df = self._sample_df()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            write_hierarchy_workbook([df], path)
            wb = load_workbook(path)
            ws = wb.active
            # Multi-row header present
            self.assertIsNotNone(ws.cell(1, 1).value)
            ranges = {str(r) for r in ws.merged_cells.ranges}
            # Group title should merge horizontally over empty child slots.
            self.assertTrue(any(":" in r for r in ranges))
            # Body STT
            # Find row with "I"
            found_i = False
            for r in range(1, 12):
                if ws.cell(r, 1).value == "I":
                    found_i = True
                    self.assertTrue(ws.cell(r, 1).font.bold)
                    break
            self.assertTrue(found_i)


if __name__ == "__main__":
    unittest.main()
