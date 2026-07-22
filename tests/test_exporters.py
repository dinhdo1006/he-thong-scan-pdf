"""Offline tests for unified plain-text / DOCX export -- no GPU, no PDF needed."""

from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd

from pdf_extractor.exporters import (
    compose_document_txt,
    dataframe_to_plaintext_table,
    export_document,
    export_document_from_markdown,
    merge_table_preserving_form,
    resolve_output_path,
    save_unified_txt,
    table_shape_similarity,
)


SAMPLE_MD = """# Report

Intro paragraph.

| Name | Age |
|------|-----|
| Alice | 30 |

More text after table.
"""

TWO_TABLES_MD = """# Report

First table:

| Name | Age |
|------|-----|
| Alice | 30 |

Middle prose.

| City | Code |
|------|------|
| Hue | 54 |

End.
"""

GARBLED_MD = """Header line.

| hoyết Tình | Erở nh | col_3 |
|------|-----|-----|
| 1 | Thu | 10.000 |

Footer.
"""


class TestComposeDocumentTxt(unittest.TestCase):
    def test_uses_extracted_table_when_available(self) -> None:
        extracted = [pd.DataFrame([["99", "88"]], columns=["Name", "Age"])]
        txt = compose_document_txt(SAMPLE_MD, extracted, apply_ocr_cleanup=False)
        self.assertIn("Intro paragraph", txt)
        self.assertIn("More text after table", txt)
        self.assertIn("99", txt)
        self.assertIn("88", txt)
        self.assertNotIn("Alice", txt)
        self.assertIn("|", txt)

    def test_falls_back_to_marker_table_when_extraction_empty(self) -> None:
        txt = compose_document_txt(SAMPLE_MD, [], apply_ocr_cleanup=False)
        self.assertIn("Alice", txt)
        self.assertIn("30", txt)
        self.assertNotIn("|---|", txt)
        self.assertIn("| Alice", txt)

    def test_preserves_reading_order(self) -> None:
        txt = compose_document_txt(SAMPLE_MD, [], apply_ocr_cleanup=False)
        intro_pos = txt.find("Intro paragraph")
        table_pos = txt.find("Alice")
        after_pos = txt.find("More text after")
        self.assertLess(intro_pos, table_pos)
        self.assertLess(table_pos, after_pos)

    def test_rejects_mismatched_extracted_table_keeps_marker_position(self) -> None:
        wrong = [pd.DataFrame([["x", "y", "z", "w"]], columns=["A", "B", "C", "D"])]
        txt = compose_document_txt(SAMPLE_MD, wrong, apply_ocr_cleanup=False)
        self.assertIn("Alice", txt)
        self.assertIn("30", txt)
        self.assertNotIn("| x ", txt)
        intro_pos = txt.find("Intro paragraph")
        table_pos = txt.find("Alice")
        after_pos = txt.find("More text after")
        self.assertLess(intro_pos, table_pos)
        self.assertLess(table_pos, after_pos)

    def test_matches_extracted_tables_by_shape_not_fifo_order(self) -> None:
        extracted = [
            pd.DataFrame([["Hue", "54"]], columns=["City", "Code"]),
            pd.DataFrame([["Bob", "40"]], columns=["Name", "Age"]),
        ]
        txt = compose_document_txt(TWO_TABLES_MD, extracted, apply_ocr_cleanup=False)
        self.assertIn("Bob", txt)
        self.assertIn("Hue", txt)
        self.assertNotIn("Alice", txt)
        bob_pos = txt.find("Bob")
        hue_pos = txt.find("Hue")
        self.assertLess(bob_pos, hue_pos)

    def test_does_not_append_unmatched_extracted_after_marker_tables(self) -> None:
        extracted = [
            pd.DataFrame([["99", "88"]], columns=["Name", "Age"]),
            pd.DataFrame([["orphan"]], columns=["Only"]),
        ]
        txt = compose_document_txt(SAMPLE_MD, extracted, apply_ocr_cleanup=False)
        self.assertIn("99", txt)
        self.assertNotIn("orphan", txt)

    def test_strips_markdown_image_placeholders(self) -> None:
        md = """Prose before.

![](_page_1_Picture_1.jpeg)

Prose after.
"""
        txt = compose_document_txt(md, [], apply_ocr_cleanup=False)
        self.assertIn("Prose before", txt)
        self.assertIn("Prose after", txt)
        self.assertNotIn("_page_1_Picture_1.jpeg", txt)
        self.assertNotIn("![](", txt)

    def test_garbled_marker_prefers_extracted_values(self) -> None:
        extracted = [pd.DataFrame([["1", "Thu an phi", "10000"]], columns=["So", "Noi dung", "Tien"])]
        txt = compose_document_txt(GARBLED_MD, extracted, apply_ocr_cleanup=False)
        self.assertIn("Thu an phi", txt)
        self.assertIn("10000", txt)
        self.assertNotIn("hoyết", txt)


class TestResolveOutputPath(unittest.TestCase):
    def test_file_path(self) -> None:
        with TemporaryDirectory() as tmp:
            p = resolve_output_path(Path(tmp) / "out.docx")
            self.assertEqual(p.name, "out.docx")

    def test_directory_gets_default_txt(self) -> None:
        with TemporaryDirectory() as tmp:
            p = resolve_output_path(Path(tmp) / "folder", default_name="output.txt")
            self.assertEqual(p.name, "output.txt")
            self.assertEqual(p.parent.name, "folder")


class TestDataframeToPlaintextTable(unittest.TestCase):
    def test_omits_generic_header_row(self) -> None:
        df = pd.DataFrame([["1", "value"]], columns=["col", "col_1"])
        txt = dataframe_to_plaintext_table(df)
        self.assertEqual(txt, "| 1 | value |")

    def test_aligns_columns_and_keeps_empty_cells(self) -> None:
        df = pd.DataFrame(
            [["1.3.2", "Thu tai san", "", "10.000"]],
            columns=["So", "Noi dung", "Ghi chu", "Tien"],
        )
        txt = dataframe_to_plaintext_table(df)
        lines = txt.splitlines()
        self.assertEqual(len(lines), 2)
        self.assertTrue(lines[0].startswith("|"))
        self.assertEqual(lines[0].count("|"), lines[1].count("|"))
        cells = [c.strip() for c in lines[1].split("|")]
        self.assertEqual(cells[3], "")
        self.assertEqual(cells[1], "1.3.2")
        self.assertEqual(cells[4], "10.000")


class TestMergeAndSimilarity(unittest.TestCase):
    def test_similarity_prefers_same_width(self) -> None:
        marker = pd.DataFrame([["a", "b"]], columns=["A", "B"])
        good = pd.DataFrame([["a", "c"]], columns=["A", "B"])
        bad = pd.DataFrame([["a", "b", "c", "d"]], columns=["A", "B", "C", "D"])
        self.assertGreater(table_shape_similarity(marker, good), table_shape_similarity(marker, bad))

    def test_merge_keeps_marker_headers_when_extracted_generic(self) -> None:
        marker = pd.DataFrame([["1", "x"]], columns=["So", "Ten"])
        extracted = pd.DataFrame([["9", "y"]], columns=["col", "col_1"])
        merged = merge_table_preserving_form(marker, extracted)
        self.assertEqual(list(merged.columns), ["So", "Ten"])
        self.assertEqual(merged.iloc[0]["So"], "9")

    def test_merge_prefers_extracted_when_marker_garbled(self) -> None:
        marker = pd.DataFrame([["1", "x"]], columns=["hoyết Tình", "Erở nh"])
        extracted = pd.DataFrame([["9", "y"]], columns=["So TT", "Noi dung"])
        merged = merge_table_preserving_form(marker, extracted)
        self.assertEqual(list(merged.columns), ["So TT", "Noi dung"])
        self.assertEqual(merged.iloc[0]["So TT"], "9")


class TestDocxExport(unittest.TestCase):
    def test_ordered_docx_has_table_with_empty_cell(self) -> None:
        from docx import Document

        md = """Intro.

| A | B | C |
|---|---|---|
| 1 |  | 3 |

Outro.
"""
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.docx"
            export_document_from_markdown(md, [], path, apply_ocr_cleanup=False)
            self.assertTrue(path.is_file())
            doc = Document(str(path))
            texts = [p.text for p in doc.paragraphs]
            self.assertTrue(any("Intro" in t for t in texts))
            self.assertTrue(any("Outro" in t for t in texts))
            self.assertEqual(len(doc.tables), 1)
            table = doc.tables[0]
            # header + 1 data row
            self.assertEqual(len(table.rows), 2)
            self.assertEqual(len(table.columns), 3)
            self.assertEqual(table.rows[1].cells[0].text, "1")
            self.assertEqual(table.rows[1].cells[1].text, "")
            self.assertEqual(table.rows[1].cells[2].text, "3")

    def test_export_document_legacy_helper(self) -> None:
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "legacy.docx"
            export_document(
                "Hello",
                [pd.DataFrame([["a", ""]], columns=["X", "Y"])],
                path,
            )
            self.assertTrue(path.is_file())


class TestSaveUnifiedTxt(unittest.TestCase):
    def test_writes_utf8_file(self) -> None:
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "result.txt"
            save_unified_txt("Xin chào\n", path)
            self.assertEqual(path.read_text(encoding="utf-8"), "Xin chào\n")


if __name__ == "__main__":
    unittest.main()
