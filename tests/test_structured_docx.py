"""Tests for structured DOCX builder (no Kafka)."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "ocr_pipeline_backup" / "ocr_pipeline" / "convert_to_docx"))

from structured_docx import (  # noqa: E402
    add_structured_tables_from_payload,
    create_document_from_pages_and_tables,
)


class TestStructuredDocx(unittest.TestCase):
    def test_structured_table_rows_cols(self) -> None:
        from docx import Document

        doc = Document()
        payload = {
            "version": 1,
            "tables": [
                {
                    "table_id": 0,
                    "columns": ["STT", "Ten", "Tien"],
                    "header_matrix": [
                        ["STT", "Ten", "Nhom"],
                        ["", "", "Tien"],
                    ],
                    "data": [["1", "A", "10"], ["2", "B", "20"]],
                    "page": 0,
                }
            ],
        }
        n = add_structured_tables_from_payload(doc, payload)
        self.assertEqual(n, 1)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.tables[0].rows), 4)  # 2 header + 2 body
        self.assertEqual(len(doc.tables[0].columns), 3)

    def test_prefers_structured_over_crude(self) -> None:
        pages = [
            {
                "page_number": 1,
                "lines": [{"line_text": "Tieu de trang"}],
                "tables": [
                    {
                        "lines": [
                            {"cells": [{"text": "crude1"}, {"text": "crude2"}]},
                        ]
                    }
                ],
            }
        ]
        payload = {
            "tables": [
                {
                    "columns": ["A", "B"],
                    "data": [["x", "y"]],
                    "header_matrix": None,
                }
            ]
        }
        doc = create_document_from_pages_and_tables(pages, payload)
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any("Tieu de trang" in t for t in texts))
        self.assertEqual(len(doc.tables), 1)
        # Structured body cell
        self.assertIn("x", doc.tables[0].rows[-1].cells[0].text)

    def test_crude_fallback_without_structured(self) -> None:
        pages = [
            {
                "lines": [{"line_text": "Hello"}],
                "tables": [
                    {"lines": [{"cells": [{"text": "c1"}, {"text": "c2"}]}]},
                ],
            }
        ]
        doc = create_document_from_pages_and_tables(pages, None)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(0, 0).text, "c1")

    def test_save_roundtrip(self) -> None:
        pages = [{"lines": [{"line_text": "P1"}], "tables": []}]
        payload = {"tables": [{"columns": ["A"], "data": [["1"]], "header_matrix": None}]}
        doc = create_document_from_pages_and_tables(pages, payload)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.docx"
            doc.save(str(path))
            self.assertTrue(path.is_file())
            self.assertGreater(path.stat().st_size, 1000)


if __name__ == "__main__":
    unittest.main()
