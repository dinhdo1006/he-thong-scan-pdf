"""
Generic hierarchical Excel export for any multi-row-header table.

Uses ``df.attrs['header_matrix']`` (list of header rows) when present to:
  - write multi-tier headers
  - auto-merge horizontal group cells (non-empty followed by empties)
  - auto-merge vertical empties under a label

No form-id / B06 template dependency.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .semantic_validator import outline_level
from .table_layout import is_group_header_label, is_outline_code, normalize_outline_token

logger = logging.getLogger(__name__)

_HELPER_COLS = {"cấp", "cap", "stt_valid", "sum_check", "stt_gap"}

_THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_HEADER_FONT = Font(bold=True, size=10)
_PARENT_FONT = Font(bold=True, size=11)
_BODY_FONT = Font(size=10)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

_LABEL_FIXES = (
    (re.compile(r"^số phải thu", re.I), "Số phải thu"),
    (re.compile(r"khân\s+cấp", re.I), "khẩn cấp"),
)


def drop_semantic_helpers(df: pd.DataFrame) -> pd.DataFrame:
    """Remove validator helper columns from a table destined for export."""
    if df is None or df.empty:
        return df
    drop = [c for c in df.columns if str(c).lower().strip() in _HELPER_COLS]
    out = df.drop(columns=drop, errors="ignore").copy()
    out.attrs.update(getattr(df, "attrs", {}))
    if len(out.columns) >= 2:
        desc = out.columns[1]
        out[desc] = [_clean_label(v) for v in out[desc].tolist()]
    return out


def _clean_label(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = " ".join(str(value).replace("\n", " ").split())
    for pat, repl in _LABEL_FIXES:
        text = pat.sub(repl, text)
    if text and text[0].islower():
        text = text[0].upper() + text[1:]
    return text


def _header_matrix_for(df: pd.DataFrame) -> list[list[str]]:
    raw = getattr(df, "attrs", {}).get("header_matrix")
    width = len(df.columns)
    if isinstance(raw, list) and raw:
        matrix = []
        for row in raw:
            padded = [str(c) if c is not None else "" for c in row] + [""] * max(0, width - len(row))
            matrix.append(padded[:width])
        return matrix
    # Fallback: single flat header row.
    return [[str(c) for c in df.columns]]


def write_hierarchy_workbook(tables: Sequence[pd.DataFrame], path: str | Path) -> Path:
    """Write tables with hierarchical headers when ``header_matrix`` is available."""
    path = Path(path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for idx, df in enumerate(tables, start=1):
        clean = drop_semantic_helpers(df)
        matrix = _header_matrix_for(clean)
        title = f"Table_{idx}"[:31]
        ws = wb.create_sheet(title=title)
        if len(matrix) >= 2:
            _write_matrix_sheet(ws, clean, matrix)
        else:
            _write_flat_sheet(ws, clean)
    if not wb.sheetnames:
        ws = wb.create_sheet("No_Tables")
        ws["A1"] = "No tables found in document."
    wb.save(path)
    logger.info("Wrote hierarchical Excel (%d sheet(s)) -> %s", len(tables), path)
    return path


def _write_matrix_sheet(ws, df: pd.DataFrame, matrix: list[list[str]]) -> None:
    n_cols = len(df.columns)
    n_hdr = len(matrix)

    for r, row in enumerate(matrix, start=1):
        for c in range(n_cols):
            val = row[c] if c < len(row) else ""
            cell = ws.cell(r, c + 1, val)
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
            cell.border = _THIN

    _apply_header_merges(ws, matrix)

    for row_i in range(len(df)):
        excel_r = n_hdr + 1 + row_i
        stt = normalize_outline_token(df.iloc[row_i, 0]) if n_cols else ""
        depth = outline_level(stt)
        is_parent = is_outline_code(stt) and "." not in stt
        for c in range(n_cols):
            val = df.iloc[row_i, c]
            text = "" if val is None or (isinstance(val, float) and pd.isna(val)) else str(val)
            if c == 1:
                text = _clean_label(text)
            cell = ws.cell(excel_r, c + 1, text)
            cell.border = _THIN
            cell.font = _PARENT_FONT if is_parent else _BODY_FONT
            if c == 0:
                cell.alignment = _CENTER
            elif c == 1:
                indent = max(0, min(depth - 1, 4))
                cell.alignment = Alignment(
                    horizontal="left", vertical="center", wrap_text=True, indent=indent
                )
            else:
                cell.alignment = Alignment(
                    horizontal="right" if text else "center",
                    vertical="center",
                    wrap_text=True,
                )

    for c in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14 if c != 2 else 42
    if n_cols >= 1:
        ws.column_dimensions["A"].width = 10
    for r in range(1, n_hdr + 1):
        ws.row_dimensions[r].height = 28 if r == 1 else 20
    ws.freeze_panes = f"A{n_hdr + 1}"


def _apply_header_merges(ws, matrix: list[list[str]]) -> None:
    """Merge empty runs under/beside labels for parent/child header layout."""
    n_rows = len(matrix)
    n_cols = max(len(r) for r in matrix) if matrix else 0

    # Horizontal merges: non-empty followed by empties — only for group-like labels
    # or when a lower header row has labels in those empty slots (parent over children).
    for r in range(n_rows):
        c = 0
        while c < n_cols:
            cell = matrix[r][c].strip() if c < len(matrix[r]) else ""
            if not cell:
                c += 1
                continue
            end = c
            while end + 1 < n_cols:
                nxt = matrix[r][end + 1].strip() if end + 1 < len(matrix[r]) else ""
                if nxt:
                    break
                end += 1
            if end > c and (
                is_group_header_label(cell)
                or _lower_rows_have_labels(matrix, r, c + 1, end)
            ):
                ws.merge_cells(
                    start_row=r + 1,
                    start_column=c + 1,
                    end_row=r + 1,
                    end_column=end + 1,
                )
            c = end + 1

    # Vertical merges: empty cells under a filled header in the same column.
    for c in range(n_cols):
        r = 0
        while r < n_rows:
            cell = matrix[r][c].strip() if c < len(matrix[r]) else ""
            if not cell:
                r += 1
                continue
            end = r
            while end + 1 < n_rows:
                below = matrix[end + 1][c].strip() if c < len(matrix[end + 1]) else ""
                if below:
                    break
                end += 1
            if end > r:
                ws.merge_cells(
                    start_row=r + 1,
                    start_column=c + 1,
                    end_row=end + 1,
                    end_column=c + 1,
                )
            r = end + 1


def _lower_rows_have_labels(
    matrix: list[list[str]], row_idx: int, start_c: int, end_c: int
) -> bool:
    """True if any row below has a non-empty label in [start_c, end_c]."""
    for r in range(row_idx + 1, len(matrix)):
        for c in range(start_c, end_c + 1):
            if c < len(matrix[r]) and matrix[r][c].strip():
                return True
    return False


def _write_flat_sheet(ws, df: pd.DataFrame) -> None:
    for c, name in enumerate(df.columns, start=1):
        cell = ws.cell(1, c, str(name))
        cell.font = _HEADER_FONT
        cell.border = _THIN
        cell.alignment = _CENTER
    for r in range(len(df)):
        for c in range(len(df.columns)):
            val = df.iloc[r, c]
            text = "" if val is None or (isinstance(val, float) and pd.isna(val)) else str(val)
            cell = ws.cell(r + 2, c + 1, text)
            cell.border = _THIN
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
