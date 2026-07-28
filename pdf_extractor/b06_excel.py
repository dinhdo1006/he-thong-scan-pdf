"""
Excel layout for B06/CD_02 forms: multi-row parent/child headers + row hierarchy.

Horizontal (header):
  Row1 group title ``Trong đó Chấp hành viên...`` merges over cols 2–5
  Row2 sub-labels Ủy thác / Trả đơn / Đình chỉ / Miễn, giảm
  Row3 column codes A B 1 2 3 4 5 6 7

Vertical (body):
  Bold parent outline rows; keep indented labels from semantic annotate.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .table_layout import (
    _B06_CANONICAL_HEADERS,
    is_outline_code,
    looks_like_b06_cd02_table,
    normalize_outline_token,
)

logger = logging.getLogger(__name__)

_HELPER_COLS = {"cấp", "cap", "stt_valid", "sum_check", "stt_gap"}

# Display names for the 3-row header (independent of DataFrame column titles).
_B06_TOP = [
    "Số TT",
    "Tiêu chí trong quyết định thi hành án",
    "Tổng số tiền, giá trị tài sản phải thi hành",
    "Trong đó Chấp hành viên đã giải quyết bằng các biện pháp",  # spans 2–5
    "",  # covered by merge
    "",
    "",
    "Số thực thu thi hành án\n(bao gồm cả số giao nhận tháng)",
    "Số đã nộp Nhà nước và chi trả đương sự\n(bao gồm cả số giao nhận tháng)",
]
_B06_MID = [
    "",
    "",
    "",
    "Ủy thác THA",
    "Trả đơn THA",
    "Đình chỉ THA",
    "Miễn, giảm THA",
    "",
    "",
]
_B06_CODE = ["A", "B", "1", "2", "3", "4", "5", "6", "7"]

_THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_HEADER_FONT = Font(bold=True, size=10)
_PARENT_FONT = Font(bold=True, size=11)
_BODY_FONT = Font(size=10)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def is_b06_excel_candidate(df: pd.DataFrame) -> bool:
    """True when the frame carries the B06 9-column schema (possibly + helpers)."""
    if df is None or df.empty:
        return False
    cols = [str(c) for c in df.columns]
    # Strip helpers and see if remaining starts with canonical names.
    form_cols = [c for c in cols if c.lower().strip() not in _HELPER_COLS]
    if len(form_cols) < 9:
        return False
    # Match by presence of key canonical headers (tolerant of Cap prefix).
    joined = " | ".join(form_cols).lower()
    hits = 0
    for key in ("số tt", "tiêu chí", "ủy thác", "trả đơn", "đình chỉ", "miễn"):
        if key in joined:
            hits += 1
    if hits >= 4:
        return True
    # Fallback: rebuild without Cap and test width-9 detector.
    slim = df[[c for c in df.columns if str(c).lower().strip() not in _HELPER_COLS]].copy()
    if len(slim.columns) == 9:
        slim.columns = list(_B06_CANONICAL_HEADERS)
        return looks_like_b06_cd02_table(slim)
    return False


def write_b06_workbook(tables: Sequence[pd.DataFrame], path: str | Path) -> Path:
    """Write one or more B06 tables with hierarchical headers into ``path``."""
    path = Path(path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    # Remove default sheet; create per table.
    default = wb.active
    wb.remove(default)

    written = 0
    for idx, df in enumerate(tables, start=1):
        if not is_b06_excel_candidate(df):
            ws = wb.create_sheet(title=f"Table_{idx}"[:31])
            _write_flat_sheet(ws, df)
        else:
            ws = wb.create_sheet(title=f"B06_Table_{idx}"[:31])
            _write_b06_sheet(ws, df)
            written += 1
    if not wb.sheetnames:
        ws = wb.create_sheet("No_Tables")
        ws["A1"] = "No tables found in document."
    wb.save(path)
    logger.info("Wrote B06-aware Excel (%d hierarchical sheet(s)) -> %s", written, path)
    return path


def _split_form_and_helpers(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    cols = [str(c) for c in df.columns]
    helpers = [c for c in cols if c.lower().strip() in _HELPER_COLS]
    # Prefer exact canonical headers in order.
    form_cols: list[str] = []
    for want in _B06_CANONICAL_HEADERS:
        if want in cols:
            form_cols.append(want)
    if len(form_cols) < 9:
        # Fallback: non-helper columns in existing order, first 9.
        form_cols = [c for c in cols if c not in helpers][:9]
    body = df[form_cols].copy()
    if len(body.columns) == 9:
        body.columns = list(_B06_CANONICAL_HEADERS)
    return body, helpers


def _write_b06_sheet(ws, df: pd.DataFrame) -> None:
    body, helpers = _split_form_and_helpers(df)
    cap_col = next((c for c in df.columns if str(c).lower().strip() in {"cấp", "cap"}), None)
    offset = 1 if cap_col is not None else 0
    n_form = len(body.columns)
    header_rows = 3

    # --- Header row 1 (group) ---
    r1 = 1
    if offset:
        ws.cell(r1, 1, "Cấp")
    for i, text in enumerate(_B06_TOP[:n_form]):
        ws.cell(r1, offset + i + 1, text)

    # --- Header row 2 (sub-labels) ---
    r2 = 2
    if offset:
        ws.cell(r2, 1, "")
    for i, text in enumerate(_B06_MID[:n_form]):
        ws.cell(r2, offset + i + 1, text)

    # --- Header row 3 (codes) ---
    r3 = 3
    if offset:
        ws.cell(r3, 1, "")
    for i, text in enumerate(_B06_CODE[:n_form]):
        ws.cell(r3, offset + i + 1, text)

    def _merge(rr1: int, cc1: int, rr2: int, cc2: int) -> None:
        ws.merge_cells(start_row=rr1, start_column=cc1, end_row=rr2, end_column=cc2)

    if offset:
        _merge(1, 1, 3, 1)
    for col in (1, 2, 3):
        if n_form >= col:
            _merge(1, offset + col, 2, offset + col)
    if n_form >= 7:
        _merge(1, offset + 4, 1, offset + 7)
    for col in (8, 9):
        if n_form >= col:
            _merge(1, offset + col, 2, offset + col)

    total_cols = offset + n_form
    for r in range(1, header_rows + 1):
        for c in range(1, total_cols + 1):
            cell = ws.cell(r, c)
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
            cell.border = _THIN

    for row_i in range(len(body)):
        excel_r = header_rows + 1 + row_i
        level = 0
        if cap_col is not None:
            raw = df.iloc[row_i][cap_col]
            try:
                level = int(raw) if pd.notna(raw) else 0
            except (TypeError, ValueError):
                level = 0
            ws.cell(excel_r, 1, level if level else "")
        stt = normalize_outline_token(body.iloc[row_i, 0]) if n_form else ""
        is_parent = level == 1 or (is_outline_code(stt) and "." not in stt)
        for c in range(n_form):
            val = body.iloc[row_i, c]
            if val is None or (isinstance(val, float) and pd.isna(val)):
                text = ""
            else:
                text = str(val)
            cell = ws.cell(excel_r, offset + c + 1, text)
            cell.border = _THIN
            cell.font = _PARENT_FONT if is_parent else _BODY_FONT
            cell.alignment = _CENTER if c == 0 else _LEFT

        if offset:
            cap_cell = ws.cell(excel_r, 1)
            cap_cell.border = _THIN
            cap_cell.alignment = _CENTER
            cap_cell.font = _PARENT_FONT if is_parent else _BODY_FONT

    helper_names = [h for h in helpers if str(h).lower().strip() not in {"cấp", "cap"}]
    for h_i, hname in enumerate(helper_names):
        col = total_cols + 1 + h_i
        ws.cell(1, col, str(hname))
        _merge(1, col, 3, col)
        for r in range(1, 4):
            cell = ws.cell(r, col)
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
            cell.border = _THIN
        for row_i in range(len(df)):
            val = df.iloc[row_i][hname]
            text = "" if val is None or (isinstance(val, float) and pd.isna(val)) else str(val)
            cell = ws.cell(header_rows + 1 + row_i, col, text)
            cell.border = _THIN
            cell.alignment = _CENTER

    widths = [6, 10, 42, 14, 12, 12, 12, 12, 14, 16]
    for i, w in enumerate(widths[:total_cols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 18
    ws.freeze_panes = "A4"


def _write_flat_sheet(ws, df: pd.DataFrame) -> None:
    """Fallback: plain header + body (non-B06 tables)."""
    for c, name in enumerate(df.columns, start=1):
        cell = ws.cell(1, c, str(name))
        cell.font = _HEADER_FONT
        cell.border = _THIN
        cell.alignment = _CENTER
    for r in range(len(df)):
        for c, name in enumerate(df.columns, start=1):
            val = df.iloc[r, c - 1]
            text = "" if val is None or (isinstance(val, float) and pd.isna(val)) else str(val)
            cell = ws.cell(r + 2, c, text)
            cell.border = _THIN
            cell.alignment = _LEFT
